#!/usr/bin/env python3
"""
Script de exportación de todas las entidades de la base de datos Ripley OLTP
a un único libro de Microsoft Excel (.xlsx), donde cada tabla es una hoja independiente.
"""

import os
import sys
from pathlib import Path
from dotenv import load_dotenv
import pandas as pd
from sqlalchemy import create_engine, inspect
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Cargar variables de entorno si existen
load_dotenv()

DB_HOST = os.getenv("MYSQL_HOST", "localhost")
DB_PORT = os.getenv("MYSQL_PORT", "3306")
DB_USER = os.getenv("MYSQL_USER", "ripley_user")
DB_PASS = os.getenv("MYSQL_PASSWORD", "ripley_pass")
DB_NAME = os.getenv("MYSQL_DATABASE", "ripley_oltp")

OUTPUT_DIR = Path("exports")
OUTPUT_FILE = OUTPUT_DIR / f"{DB_NAME}_data.xlsx"


def main():
    print("=" * 60)
    print(f"📊 EXPORTADOR DE ENTIDADES OLTP A EXCEL - RIPLEY")
    print(f"🔗 Conectando a {DB_USER}@{DB_HOST}:{DB_PORT}/{DB_NAME}...")
    print("=" * 60)

    connection_url = (
        f"mysql+pymysql://{DB_USER}:{DB_PASS}@{DB_HOST}:{DB_PORT}/{DB_NAME}?charset=utf8mb4"
    )

    try:
        engine = create_engine(connection_url)
        inspector = inspect(engine)
        tables = inspector.get_table_names()
    except Exception as e:
        print(f"\n❌ Error al conectar a la base de datos MySQL:\n   {e}")
        print("\n💡 Verifica que el contenedor esté corriendo con: make db-up (o sudo docker compose up -d)")
        sys.exit(1)

    if not tables:
        print("⚠️ No se encontraron tablas en la base de datos.")
        sys.exit(0)

    # Asegurar directorio de salida
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print(f"\n📁 Se encontraron {len(tables)} tablas. Procesando...")

    summary = []

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        for table in tables:
            # Excel limita el nombre de las hojas a 31 caracteres
            sheet_name = table[:31]

            try:
                df = pd.read_sql_table(table, con=engine)
            except Exception as e:
                print(f"⚠️ Error leyendo tabla {table}: {e}")
                continue

            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]

            # Estilo para la fila de encabezados
            header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style="thin", color="D9D9D9"),
                right=Side(style="thin", color="D9D9D9"),
                top=Side(style="thin", color="D9D9D9"),
                bottom=Side(style="thin", color="D9D9D9"),
            )

            for col_idx, col in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

                # Auto-ajuste de ancho de columna
                max_len = max(
                    len(str(col)),
                    *(len(str(val)) for val in df[col].astype(str).tolist()[:50] if val)
                ) if not df.empty else len(str(col))

                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

            # Bordes ligeros para las celdas de datos
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                for cell in row:
                    cell.border = thin_border

            summary.append({"Tabla": table, "Registros": len(df), "Hoja Excel": sheet_name})
            print(f"  ✓ [{len(df):>4} filas] {table}")

    print("\n" + "=" * 60)
    print(f"✅ EXPORTACIÓN EXITOSA")
    print(f"📄 Archivo generado: {OUTPUT_FILE.resolve()}")
    print(f"📑 Total de entidades exportadas: {len(summary)}")
    print("=" * 60)


if __name__ == "__main__":
    main()
